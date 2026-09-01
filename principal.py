import os
import re
import sys
import time
import numpy as np
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
import pandas as pd
from rapidfuzz import fuzz, process
import warnings
warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')

pd.set_option('future.no_silent_downcasting', True)

# ------------------------------------------------------------------------------
# FUNÇÕES DE NORMALIZAÇÃO E LIMPEZA
# ------------------------------------------------------------------------------
def normalizar_texto_simples(texto):
    """Normaliza uma string individual de forma rápida sem instanciar Series Pandas."""
    if not isinstance(texto, str):
        return ""
    texto = texto.lower().strip()
    texto = re.sub(r"[^\w\s]", "", texto)
    texto = re.sub(r"\s+", " ", texto)
    return texto.strip()


def normalizar(s):
    """Normaliza uma Série de strings do Pandas."""
    s = s.astype(str).str.lower().str.strip()
    s = s.str.replace(r"[^\w\s]", "", regex=True)
    s = s.str.replace(r"\s+", " ", regex=True)
    return s.str.strip()


def converter_para_float(coluna):
    """Converte valores com formato de moeda brasileiro (1.000,00) para float com segurança,
    lidando perfeitamente com colunas de tipos mistos."""
    
    def limpar_moeda(valor):
        if pd.isna(valor):
            return 0.0
        
        # Se for texto, remove os pontos de milhar e troca vírgula por ponto
        if isinstance(valor, str):
            valor = valor.strip().replace(".", "").replace(",", ".")
            
        # Converte para float (funciona tanto para as strings limpas quanto para números pré-existentes)
        return float(valor)

    return coluna.apply(limpar_moeda)


def limpeza(df, planilha="transito"):
    if planilha == "transito":
        df = df.dropna(thresh=6)
        df = df.dropna(axis=1, thresh=5)
        df = df[df["Mercadoria"] != "Mercadoria"]

        df["Mercadoria"] = df["Mercadoria"].astype(str).str.upper()
        df["Tomador"] = df["Tomador"].astype(str).str.upper()

        df["Vlr NF"] = converter_para_float(df["Vlr NF"])
        df["Valor Frete"] = converter_para_float(df["Valor Frete"])

        combustiveis = [
            "ALCOOL",
            "BIODIESEL",
            "ETANOL",
            "GASOLINA",
            "ONU1170",
            "ONU 1170",
        ]
        padrao = "|".join(combustiveis)
        df["Operação"] = np.where(
            df["Mercadoria"].str.contains(padrao, na=False),
            "COMBUSTIVEL",
            "VEGETAL",
        )

    else:
        df = df.dropna(axis=1, thresh=5)
        df = df.dropna()
        df = df.rename(columns={"Unnamed: 2": "Empresa", "Unnamed: 18": "Qtd"})

        df["Empresa"] = df["Empresa"].astype(str).str.replace(
            "Razão Social : ", ""
        )
        df["Empresa"] = df["Empresa"].str.upper()

        df["Receber"] = converter_para_float(df["Receber"])
        df["Pagar"] = converter_para_float(df["Pagar"])

    return df


# ------------------------------------------------------------------------------
# FUZZY MATCHING E CLASSIFICAÇÃO
# ------------------------------------------------------------------------------
def fuzzymatch(df, similaridade=85):
    df["Tomador_norm"] = normalizar(df["Tomador"])
    counts = df["Tomador_norm"].value_counts()
    mapping, representatives = {}, []

    for name in counts.index:
        if representatives:
            best, score, _ = process.extractOne(
                name, representatives, scorer=fuzz.token_sort_ratio
            )
            if score >= similaridade:
                mapping[name] = best
                continue
        representatives.append(name)
        mapping[name] = name

    df["Empresa"] = df["Tomador_norm"].map(mapping).str.upper()
    return df


def fuzzymatch_cruzado(
    df_origem,
    df_destino,
    col_origem="Empresa",
    col_destino="Empresa",
    similaridade=85,
):
    referencias_norm = df_origem[col_origem].dropna().unique()
    mapa_referencias = {}
    for ref in referencias_norm:
        chave = normalizar_texto_simples(ref)
        if chave not in mapa_referencias:
            mapa_referencias[chave] = ref

    chaves_referencia = list(mapa_referencias.keys())

    mapping_resultados = {}
    empresas_destino = df_destino[col_destino].dropna().unique()

    for emp_raw in empresas_destino:
        emp_norm = normalizar_texto_simples(emp_raw)

        if emp_norm in mapa_referencias:
            mapping_resultados[emp_raw] = {
                "Empresa_Padronizada": mapa_referencias[emp_norm],
                "Score": 100,
                "Status_Match": "Exato",
            }
        elif chaves_referencia:
            best_match, score, _ = process.extractOne(
                emp_norm, chaves_referencia, scorer=fuzz.token_sort_ratio
            )
            if score >= similaridade:
                mapping_resultados[emp_raw] = {
                    "Empresa_Padronizada": mapa_referencias[best_match],
                    "Score": round(score, 1),
                    "Status_Match": "Fuzzy Match",
                }
            else:
                mapping_resultados[emp_raw] = {
                    "Empresa_Padronizada": emp_raw,
                    "Score": round(score, 1),
                    "Status_Match": "Sem Match (Novo)",
                }
        else:
            mapping_resultados[emp_raw] = {
                "Empresa_Padronizada": emp_raw,
                "Score": 0,
                "Status_Match": "Sem Match (Novo)",
            }

    df_destino["Empresa_Padronizada"] = df_destino[col_destino].map(
        lambda x: mapping_resultados.get(x, {}).get("Empresa_Padronizada", x)
    )
    return df_destino, mapping_resultados


def gerenciar_classificacao_manual(
    df_transito,
    df_descarregados,
    arquivo_pendentes="Classificar_Pendentes.xlsx",
):
    mapa_transito = (
        df_transito.groupby("Empresa_Padronizada")["Operação"].last().to_dict()
    )

    mapa_manual = {}
    if os.path.exists(arquivo_pendentes):
        try:
            df_pendentes_existente = pd.read_excel(arquivo_pendentes)
            for _, row in df_pendentes_existente.iterrows():
                emp = str(row["Empresa"]).strip()
                cat = str(row["Categoria"]).strip()
                if cat in ["VEGETAL", "COMBUSTIVEL", "OUTROS"]:
                    mapa_manual[emp] = cat
        except Exception as e:
            print(f"Aviso: Não foi possível ler o histórico de pendências: {e}")

    registros_pendentes = []
    mapa_consolidado = {}
    empresas_descarregados = df_descarregados[
        "Empresa_Padronizada"
    ].unique()

    for emp_padronizada in empresas_descarregados:
        if emp_padronizada in mapa_transito:
            categoria_final = mapa_transito[emp_padronizada]
        elif emp_padronizada in mapa_manual:
            categoria_final = mapa_manual[emp_padronizada]
        else:
            categoria_final = "OUTROS"

        mapa_consolidado[emp_padronizada] = categoria_final
        registros_pendentes.append(
            {"Empresa": emp_padronizada, "Categoria": categoria_final}
        )

    df_pendentes = pd.DataFrame(registros_pendentes)
    df_pendentes["Ordem"] = df_pendentes["Categoria"].apply(
        lambda x: 0 if x == "OUTROS" else 1
    )
    df_pendentes = df_pendentes.sort_values(by=["Ordem", "Empresa"]).drop(
        columns=["Ordem"]
    )

    try:
        df_pendentes.to_excel(arquivo_pendentes, index=False)
        wb = openpyxl.load_workbook(arquivo_pendentes)
        ws = wb.active

        dv = DataValidation(
            type="list",
            formula1='"VEGETAL,COMBUSTIVEL,OUTROS"',
            allow_blank=True,
        )
        dv.error = (
            "Escolha uma das opções da lista: VEGETAL, COMBUSTIVEL ou OUTROS"
        )
        dv.errorTitle = "Opção Inválida"
        dv.prompt = "Selecione a categoria da empresa"
        dv.promptTitle = "Categoria"

        ws.add_data_validation(dv)
        max_row = len(df_pendentes) + 1
        if max_row >= 2:
            dv.add(f"B2:B{max_row}")

        ws.column_dimensions["A"].width = 65
        ws.column_dimensions["B"].width = 20
        wb.save(arquivo_pendentes)
        wb.close()
    except PermissionError:
        print(
            f"\n[ERRO] Feche o arquivo '{arquivo_pendentes}' antes de continuar."
        )
        raise

    return mapa_consolidado


def normalizar_avancado(texto):
    """Remove pontuações isoladas, pontuações entre letras e padroniza os espaços."""
    if not isinstance(texto, str):
        return ""
    texto = texto.upper().strip()
    texto = re.sub(r"[.,;\-/]", " ", texto)
    texto = re.sub(r"\s+", " ", texto)
    return texto.strip()


def fuzzymatch_consolidação(
    df, col_empresa="Empresa", similaridade_corte=85
):
    """Aplica o agrupamento fuzzy nos dados consolidados, unificando variações
    sutis de nomes de empresas."""
    
    df_copia = df.copy()
    empresas_unicas = df_copia[col_empresa].dropna().unique()
    nomes_limpos = {emp: normalizar_avancado(emp) for emp in empresas_unicas}

    mapeamento_canonico = {}
    representantes_limpos = []
    mapa_limpo_para_original = {}

    for emp_orig in empresas_unicas:
        emp_limpa = nomes_limpos[emp_orig]
        if not emp_limpa:
            continue

        if representantes_limpos:
            best_match, score, _ = process.extractOne(
                emp_limpa,
                representantes_limpos,
                scorer=fuzz.token_set_ratio,
            )
            if score >= similaridade_corte:
                nome_canonico = mapa_limpo_para_original[best_match]
                mapeamento_canonico[emp_orig] = nome_canonico
                continue

        representantes_limpos.append(emp_limpa)
        mapa_limpo_para_original[emp_limpa] = emp_orig
        mapeamento_canonico[emp_orig] = emp_orig

    df_copia[col_empresa] = df_copia[col_empresa].map(mapeamento_canonico)
    return df_copia


def gerar_pivot_categoria(df_descarregados, df_transito, categoria):
    """ Gera a tabela consolidada (Descarregados x Em Trânsito) por empresa
    para a categoria informada ('VEGETAL', 'COMBUSTIVEL' ou 'OUTROS')."""

    sub_des = df_descarregados[df_descarregados["Operação"] == categoria]
    sub_tra = df_transito[df_transito["Operação"] == categoria]

    if not sub_des.empty:
        p_des = sub_des.pivot_table(
            index="Empresa_Padronizada", values="Receber", aggfunc="sum"
        ).rename(columns={"Receber": "Descarregados"})
    else:
        # Adicionado dtype=float para evitar criar a coluna como object
        p_des = pd.DataFrame(columns=["Descarregados"], dtype=float)
        p_des.index.name = "Empresa_Padronizada"

    if not sub_tra.empty:
        p_tra = sub_tra.pivot_table(
            index="Empresa_Padronizada", values="Valor Frete", aggfunc="sum"
        ).rename(columns={"Valor Frete": "Em Trânsito"})
    else:
        # Adicionado dtype=float para evitar criar a coluna como object
        p_tra = pd.DataFrame(columns=["Em Trânsito"], dtype=float)
        p_tra.index.name = "Empresa_Padronizada"

    resumo = pd.merge(p_des, p_tra, on="Empresa_Padronizada", how="outer")
    
    # Com as colunas nascendo como float, o fillna(0.0) não fará downcasting
    resumo = resumo.fillna(0.0)
    
    resumo = resumo.reset_index().rename(
        columns={"Empresa_Padronizada": "Empresa"}
    )
    return resumo


# ------------------------------------------------------------------------------
# EXECUÇÃO PRINCIPAL
# ------------------------------------------------------------------------------
def main():
    print("Iniciando o processamento das planilhas...")

    # 1. Resolução absoluta de diretórios ANTES do try block
    if getattr(sys, 'frozen', False):
        diretorio_base = os.path.dirname(sys.executable)
    else:
        diretorio_base = os.path.dirname(os.path.abspath(__file__))
    
    # Declarando os caminhos das planilhas usando o diretório do executável
    caminho_transito = os.path.join(diretorio_base, 'transito.xlsx')
    caminho_descarregados = os.path.join(diretorio_base, 'descarregados_cnpj.xlsx')
    caminho_pendentes = os.path.join(diretorio_base, "Classificar_Pendentes.xlsx")

    # 2. Uso dos caminhos declarados
    try:
        df_t = pd.read_excel(caminho_transito, sheet_name="ReportXML", header=9)
        des = pd.read_excel(caminho_descarregados, header=8)
    except FileNotFoundError as e:
        print(f"\n[ERRO] Arquivo de entrada não encontrado: {e.filename}")
        print(f"Certifique-se de que os arquivos Excel estão na pasta:\n{diretorio_base}")
        return

    colunas_obrigatorias = ["Mercadoria", "Tomador", "Vlr NF"]
    faltantes = [col for col in colunas_obrigatorias if col not in df_t.columns]
    if faltantes:
        print(f"[ERRO] A planilha 'transito.xlsx' não possui as colunas: {faltantes}")
        return
    
    # Limpeza
    df_t = limpeza(df_t)
    des = limpeza(des, planilha="descarregados")

    # Padronização de Nomes
    df_t = fuzzymatch(df_t)
    df_t["Empresa_Padronizada"] = df_t["Empresa"]

    des, _ = fuzzymatch_cruzado(
        df_origem=df_t,
        df_destino=des,
        col_origem="Empresa",
        col_destino="Empresa",
    )

    # Passando a variável de caminho para a função de gerenciamento manual
    mapa_categorias = gerenciar_classificacao_manual(df_t, des, arquivo_pendentes=caminho_pendentes)
    des["Operação"] = des["Empresa_Padronizada"].map(mapa_categorias)

    # Consolidação
    p_c = gerar_pivot_categoria(des, df_t, "COMBUSTIVEL")
    p_v = gerar_pivot_categoria(des, df_t, "VEGETAL")
    p_o = gerar_pivot_categoria(des, df_t, "OUTROS")

    # Fuzzy na Consolidação
    p_c = fuzzymatch_consolidação(
        p_c, col_empresa="Empresa", similaridade_corte=85
    )
    p_v = fuzzymatch_consolidação(
        p_v, col_empresa="Empresa", similaridade_corte=85
    )
    p_o = fuzzymatch_consolidação(
        p_o, col_empresa="Empresa", similaridade_corte=85
    )

    p_c = p_c.groupby("Empresa", as_index=False).agg(
        {"Descarregados": "sum", "Em Trânsito": "sum"}
    )
    p_v = p_v.groupby("Empresa", as_index=False).agg(
        {"Descarregados": "sum", "Em Trânsito": "sum"}
    )
    p_o = p_o.groupby("Empresa", as_index=False).agg(
        {"Descarregados": "sum", "Em Trânsito": "sum"}
    )

    # ==============================================================================
    # GRAVAÇÃO NA PLANILHA FINAL VIA XLSXWRITER
    # ==============================================================================
    data_formatada = time.strftime("%d-%m-%Y")
    
    # 3. Aplicar também ao arquivo final de saída
    nome_saida = os.path.join(diretorio_base, f"{data_formatada}_final.xlsx")

    try:
        with pd.ExcelWriter(nome_saida, engine="xlsxwriter") as writer:
            workbook = writer.book
            moeda = workbook.add_format({"num_format": "R$ #,##0.00"})
            moeda_total = workbook.add_format(
                {"num_format": "R$ #,##0.00", "bold": True}
            )

            # -------------------------------------------------------------
            # 1. ABA: Base Trânsito
            # -------------------------------------------------------------
            df_t_base = df_t[
                [
                    "Dt. Emissão",
                    "Operação",
                    "Mercadoria",
                    "Tomador",
                    "Valor Frete",
                    "Vlr NF",
                ]
            ]
            df_t_base.to_excel(
                writer,
                sheet_name="Base Trânsito",
                index=False,
                startrow=1,
                header=False,
            )
            ws_base = writer.sheets["Base Trânsito"]
            max_row_base = len(df_t_base) + 1
            max_col_base = len(df_t_base.columns) - 1

            ws_base.add_table(
                0,
                0,
                max_row_base,
                max_col_base,
                {
                    "name": "Tabela_Base_Transito",
                    "columns": [
                        {"header": "Dt. Emissão", "total_string": "Total"},
                        {"header": "Operação"},
                        {"header": "Mercadoria"},
                        {"header": "Tomador"},
                        {
                            "header": "Valor Frete",
                            "total_function": "sum",
                            "format": moeda,
                        },
                        {
                            "header": "Vlr NF",
                            "total_function": "sum",
                            "format": moeda,
                        },
                    ],
                    "total_row": True,
                    "style": "Table Style Light 9",
                },
            )

            ws_base.set_column("A:B", 15)
            ws_base.set_column("C:C", 20)
            ws_base.set_column("D:D", 30)
            ws_base.set_column("E:F", 15, moeda)
            ws_base.write(
                max_row_base, 4, f"=SUM(E2:E{max_row_base})", moeda_total
            )
            ws_base.write(
                max_row_base, 5, f"=SUM(F2:F{max_row_base})", moeda_total
            )

            # -------------------------------------------------------------
            # 2. ABA: Base Descarregados
            # -------------------------------------------------------------
            des_base = des.copy()
            des_base.to_excel(
                writer,
                sheet_name="Base Descarregados",
                index=False,
                startrow=1,
                header=False,
            )
            ws_des = writer.sheets["Base Descarregados"]
            max_row_des = len(des_base) + 1
            max_col_des = len(des_base.columns) - 1

            colunas_des = []
            for col_name in des_base.columns:
                if col_name in ["Receber", "Pagar"]:
                    colunas_des.append(
                        {
                            "header": col_name,
                            "total_function": "sum",
                            "format": moeda,
                        }
                    )
                else:
                    colunas_des.append({"header": col_name})

            ws_des.add_table(
                0,
                0,
                max_row_des,
                max_col_des,
                {
                    "name": "Tabela_Base_Descarregados",
                    "columns": colunas_des,
                    "total_row": True,
                    "style": "Table Style Light 9",
                },
            )

            ws_des.set_column("A:A", 45)
            ws_des.set_column("B:B", 15)
            ws_des.set_column("C:D", 5)
            ws_des.set_column("E:E", 45)
            ws_des.set_column("F:F", 20)

            # -------------------------------------------------------------
            # 3. ABAS: Resumo (Combustível, Vegetal, Outros)
            # -------------------------------------------------------------
            abas_config = [
                ("Combustível", p_c, "Tabela_Combustivel"),
                ("Vegetal", p_v, "Tabela_Vegetal"),
                ("Outros", p_o, "Tabela_Outros"),
            ]

            for nome_aba, df_pivot, nome_tabela in abas_config:
                df_pivot.to_excel(
                    writer,
                    sheet_name=nome_aba,
                    index=False,
                    startrow=1,
                    header=False,
                )
                ws = writer.sheets[nome_aba]
                max_row = len(df_pivot) + 1
                max_col = len(df_pivot.columns) - 1

                if len(df_pivot) > 0:
                    ws.add_table(
                        0,
                        0,
                        max_row,
                        max_col,
                        {
                            "name": nome_tabela,
                            "columns": [
                                {"header": "Empresa", "total_string": "Total"},
                                {
                                    "header": "Descarregados",
                                    "total_function": "sum",
                                    "format": moeda,
                                },
                                {
                                    "header": "Em Trânsito",
                                    "total_function": "sum",
                                    "format": moeda,
                                },
                            ],
                            "total_row": True,
                            "style": "Table Style Light 9",
                        },
                    )

                    ws.set_column("A:A", 60)
                    ws.set_column("B:C", 18, moeda)
                    ws.write(max_row, 1, f"=SUM(B2:B{max_row})", moeda_total)
                    ws.write(max_row, 2, f"=SUM(C2:C{max_row})", moeda_total)
                else:
                    ws.write(
                        0, 0, "Nenhum registro encontrado para esta categoria."
                    )

        print(f"\n[SUCESSO] Processamento concluído! Arquivo: '{nome_saida}'")

    except PermissionError:
        print(
            f"\n[ERRO] O arquivo '{nome_saida}' está aberto no Excel. Feche-o e tente novamente."
        )

if __name__ == "__main__":
    try:
        main()
    except Exception as err:
        print(f"\n[ERRO INESPERADO]: {err}")
    finally:
        input("\nPressione ENTER para encerrar...")
